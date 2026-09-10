"""
Documents: extraction, page rendering, and joining the timeline to disk.

Two facts drive this module.

1. The `documents/` folder is a **subset** of what the timeline says was
   uploaded, stored under cleaned-up filenames (`074038651 Notice of Related
   Cases.pdf` on the timeline is `notice_of_related_cases_filed.pdf` on disk).
   So the join is fuzzy, and "uploaded but not in this snapshot" is a real,
   reportable state rather than an error.

2. Some PDFs have no text layer at all (six in case_001 — the issued summons,
   the return of service, the standing family law order). They are detected
   here and OCR'd once at ingest; nothing downstream needs to special-case
   them.
"""

from __future__ import annotations

import io
import re
from pathlib import Path

import pypdfium2 as pdfium
from pydantic import BaseModel, Field
from rapidfuzz import fuzz

from ..data_loader import list_documents
from .models import NormalizedEvent

# A page with fewer extractable characters than this is treated as an image.
# Text-layer pages in these cases run 800–3000 chars; scanned pages return 0.
_SCAN_CHAR_THRESHOLD = 40

# Fuzzy score above which a timeline filename and a disk filename are the same
# document. Tuned against the real pairs in both cases: every true pair
# scores 100 and the closest false pair scores 62.
_MATCH_THRESHOLD = 70

_SUPPORTED = {".pdf", ".docx"}


class PageText(BaseModel):
    page: int
    text: str


class DocumentText(BaseModel):
    name: str
    kind: str
    pages: list[PageText] = Field(default_factory=list)
    is_scanned: bool = False
    ocr_applied: bool = False

    @property
    def char_count(self) -> int:
        return sum(len(p.text) for p in self.pages)


class CatalogEntry(BaseModel):
    """One document, as seen from both the timeline and the filesystem."""

    display_name: str
    disk_name: str | None = None
    timeline_file_name: str | None = None
    category: str | None = None
    who_uploaded: str | None = None
    summary: str | None = None
    uploaded_at: str | None = None
    event_refs: list[str] = Field(default_factory=list)
    page_count: int | None = None
    # "text": extractable; "scanned": image-only, OCR needed; "absent": the
    # timeline records an upload but the file is not in this snapshot;
    # "unsupported": on disk but not a format we extract.
    text_status: str = "absent"

    @property
    def readable(self) -> bool:
        return self.disk_name is not None and self.text_status in ("text", "scanned")


# --------------------------------------------------------------------------
# Extraction
# --------------------------------------------------------------------------


def extract_docx(path: Path) -> DocumentText:
    """DOCX has no pages; the whole body is returned as page 1.

    Tables carry real content in these filings (party details, case captions),
    so they are flattened into the text rather than skipped.
    """
    import docx  # imported lazily; only needed when a .docx is actually read

    document = docx.Document(str(path))
    parts: list[str] = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            # Word repeats merged cells; collapse consecutive duplicates.
            deduped = [c for i, c in enumerate(cells) if i == 0 or c != cells[i - 1]]
            if deduped:
                parts.append(" | ".join(deduped))

    text = re.sub(r"\n{3,}", "\n\n", "\n".join(parts)).strip()
    return DocumentText(name=path.name, kind="docx", pages=[PageText(page=1, text=text)])


def extract_pdf(path: Path) -> DocumentText:
    """Extract the text layer page by page, flagging image-only documents."""
    pdf = pdfium.PdfDocument(str(path))
    try:
        pages = []
        for number in range(len(pdf)):
            page = pdf[number]
            textpage = page.get_textpage()
            raw = textpage.get_text_range() or ""
            pages.append(PageText(page=number + 1, text=raw.replace("\r\n", "\n").strip()))
    finally:
        pdf.close()

    scanned = all(len(p.text) < _SCAN_CHAR_THRESHOLD for p in pages)
    return DocumentText(name=path.name, kind="pdf", pages=pages, is_scanned=scanned)


def extract(path: Path) -> DocumentText:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return extract_docx(path)
    if suffix == ".pdf":
        return extract_pdf(path)
    return DocumentText(name=path.name, kind=suffix.lstrip(".") or "unknown")


def page_count(path: Path) -> int | None:
    if path.suffix.lower() != ".pdf":
        return 1 if path.suffix.lower() == ".docx" else None
    pdf = pdfium.PdfDocument(str(path))
    try:
        return len(pdf)
    finally:
        pdf.close()


def render_page(path: Path, page: int, scale: float = 2.0) -> bytes:
    """Rasterize one 1-indexed PDF page to PNG.

    Used both to feed scanned pages to vision OCR at ingest and to show the
    literal page behind a citation in the UI.
    """
    pdf = pdfium.PdfDocument(str(path))
    try:
        if not 1 <= page <= len(pdf):
            raise IndexError(f"page {page} out of range for {path.name} ({len(pdf)} pages)")
        image = pdf[page - 1].render(scale=scale).to_pil()
    finally:
        pdf.close()
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


# --------------------------------------------------------------------------
# Joining the timeline to disk
# --------------------------------------------------------------------------

# Workflow words that appear on one side of a pair but not the other.
# Digits are deliberately NOT stripped: `FL-150` and `FL-140` are different
# documents, and removing numbers collapses every California form into "fl".
_NOISE = re.compile(
    r"\b(initial|draft|copy|final|signed|filed|document|documents|pdf|docx|the|and|of|for)\b",
    re.I,
)

# E-filing job numbers (`074038651 UCCJEA.pdf`, `147533756 Document...`) are
# transaction ids, not document identity. Form numbers are 3 digits or fewer.
_JOB_NUMBER = re.compile(r"\b\d{4,}\b")

# `FL-150`, `FL 150`, `fl_150` all mean the same form; fold to one token so
# the number cannot be split away from its prefix by tokenization.
_FORM_CODE = re.compile(r"\b([a-z]{2})[\s\-_]?(\d{2,3}[a-z]?)\b", re.I)


def _fingerprint(name: str) -> str:
    """Reduce a filename to the words that identify the document."""
    stem = Path(name).stem
    stem = re.sub(r"[_\-.()]+", " ", stem)
    stem = _JOB_NUMBER.sub(" ", stem)
    stem = _FORM_CODE.sub(lambda m: f"{m.group(1).lower()}{m.group(2).lower()}", stem)
    stem = _NOISE.sub(" ", stem)
    return " ".join(stem.lower().split())


def _score(left: str, right: str) -> float:
    """Similarity between two filenames.

    A form code shared by both sides is decisive, and a form code present on
    both sides but *different* is disqualifying — otherwise `FL-150` and
    `FL-140` score ~95 on their surrounding words alone.
    """
    left_fp, right_fp = _fingerprint(left), _fingerprint(right)
    if not left_fp or not right_fp:
        return 0.0

    left_codes = set(re.findall(r"\b[a-z]{2}\d{2,3}[a-z]?\b", left_fp))
    right_codes = set(re.findall(r"\b[a-z]{2}\d{2,3}[a-z]?\b", right_fp))
    if left_codes and right_codes:
        return 100.0 if left_codes & right_codes else 0.0

    # token_sort_ratio, not token_set_ratio: the latter ignores tokens that
    # only appear on one side, so a shared client-name prefix
    # (`Idris Sorlie - ...`) would make every document in the case look
    # alike. Sorting compares the full token sequence instead.
    return float(fuzz.token_sort_ratio(left_fp, right_fp))


def _assign(sources: list[str], candidates: list[str]) -> dict[int, str]:
    """Match timeline filenames to disk filenames, best pairs first.

    Global rather than per-source greedy: scoring every pair and consuming
    them in descending order stops an early weak match from stealing a file
    that a later source matches far better.
    """
    scored = [
        (_score(source, candidate), index, candidate)
        for index, source in enumerate(sources)
        for candidate in candidates
    ]
    scored.sort(key=lambda item: (-item[0], item[1], item[2]))

    assignment: dict[int, str] = {}
    used: set[str] = set()
    for score, index, candidate in scored:
        if score < _MATCH_THRESHOLD:
            break
        if index in assignment or candidate in used:
            continue
        assignment[index] = candidate
        used.add(candidate)
    return assignment


def build_catalog(case_dir: str | Path, events: list[NormalizedEvent]) -> list[CatalogEntry]:
    """Merge timeline upload events with the files actually present on disk.

    Every disk file appears exactly once. Timeline uploads that cannot be
    matched to a file are still listed, marked `absent`, so the agent can say
    "the client uploaded it, but it isn't in this snapshot" instead of either
    hiding it or pretending to read it.
    """
    case_dir = Path(case_dir)
    unclaimed = list_documents(case_dir)

    # The same file is often uploaded more than once (case_002 re-drafts the
    # whole FL-100 family twice). Group by filename first so one document is
    # one catalog entry that cites every upload event.
    grouped: dict[str, list[NormalizedEvent]] = {}
    for event in events:
        if event.doc_ref is None:
            continue
        key = (event.doc_ref.file_name or event.doc_ref.display_name).strip().lower()
        grouped.setdefault(key, []).append(event)

    source_keys = list(grouped)
    source_names = [
        (grouped[k][0].doc_ref.file_name or grouped[k][0].doc_ref.display_name)  # type: ignore[union-attr]
        for k in source_keys
    ]
    assignment = _assign(source_names, unclaimed)

    entries: list[CatalogEntry] = []
    for index, key in enumerate(source_keys):
        group = grouped[key]
        first, last = group[0], group[-1]
        ref = last.doc_ref
        assert ref is not None
        match = assignment.get(index)
        if match:
            unclaimed.remove(match)
        entries.append(
            CatalogEntry(
                display_name=ref.display_name,
                disk_name=match,
                timeline_file_name=ref.file_name,
                category=ref.category,
                who_uploaded=ref.who_uploaded,
                summary=next((e.doc_ref.summary for e in group if e.doc_ref.summary), None),
                uploaded_at=first.when.isoformat(),
                event_refs=[e.ref for e in group],
            )
        )

    # Files on disk that no upload event claimed still belong in the catalog.
    for name in unclaimed:
        entries.append(
            CatalogEntry(display_name=Path(name).stem.replace("_", " ").title(), disk_name=name)
        )

    for entry in entries:
        if not entry.disk_name:
            continue
        path = case_dir / "documents" / entry.disk_name
        suffix = path.suffix.lower()
        if suffix not in _SUPPORTED:
            entry.text_status = "unsupported"
            continue
        try:
            entry.page_count = page_count(path)
            if suffix == ".pdf":
                extracted = extract_pdf(path)
                entry.text_status = "scanned" if extracted.is_scanned else "text"
            else:
                entry.text_status = "text"
        except Exception:
            entry.text_status = "unsupported"

    entries.sort(key=lambda e: (e.disk_name is None, e.display_name.lower()))
    return entries
